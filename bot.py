"""
Main bot module for the Telegram Task Manager Bot.
Handles all Telegram interactions, commands, and conversation handlers.
"""

import os
import io
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ConversationHandler,
    filters,
    ContextTypes,
)

from config import (
    BOT_TOKEN,
    MESSAGES,
    STATUS_EMOJIS,
    REMINDER_KEYBOARD,
    REMINDER_OPTIONS,
    EDIT_FIELD_KEYBOARD,
    DAILY_SUMMARY_HOUR,
    DAILY_SUMMARY_MINUTE,
    TASK_TITLE,
    TASK_DESCRIPTION,
    TASK_DUE_DATE,
    TASK_DUE_TIME,
    TASK_REMINDER,
    EDIT_SELECT_TASK,
    EDIT_SELECT_FIELD,
    EDIT_TITLE,
    EDIT_DESCRIPTION,
    EDIT_DUE_DATE,
    EDIT_DUE_TIME,
    EDIT_REMINDER,
    LOCAL_TIMEZONE,
    DATABASE_NAME,
    logger,
    format_iran_time,
)
from database import db
from scheduler import scheduler


# -----------------------------------------------------------------------------
# Main Menu Keyboard
# -----------------------------------------------------------------------------
def get_main_menu_keyboard() -> InlineKeyboardMarkup:
    """Create the main menu inline keyboard."""
    keyboard = [
        [
            InlineKeyboardButton("📋 لیست وظایف", callback_data="main_tasks"),
            InlineKeyboardButton("➕ وظیفه جدید", callback_data="main_addtask"),
        ],
        [
            InlineKeyboardButton("📅 امروز", callback_data="main_today"),
            InlineKeyboardButton("⚠️ عقب‌افتاده", callback_data="main_overdue"),
        ],
        [
            InlineKeyboardButton("✅ تکمیل", callback_data="main_complete"),
            InlineKeyboardButton("✏️ ویرایش", callback_data="main_edittask"),
        ],
        [
            InlineKeyboardButton("🗑 حذف", callback_data="main_deletetask"),
            InlineKeyboardButton("📊 خلاصه", callback_data="main_summary"),
        ],
        [
            InlineKeyboardButton("📤 خروجی اکسل", callback_data="main_export"),
            InlineKeyboardButton("📥 ورودی اکسل", callback_data="main_import"),
        ],
        [
            InlineKeyboardButton("❓ راهنما", callback_data="main_help"),
        ],
    ]
    return InlineKeyboardMarkup(keyboard)


def get_back_button() -> InlineKeyboardMarkup:
    """Simple back to menu button."""
    return InlineKeyboardMarkup([[
        InlineKeyboardButton("🏠 منوی اصلی", callback_data="main_menu")
    ]])


# -----------------------------------------------------------------------------
# Menu Handlers (non-conversation)
# -----------------------------------------------------------------------------
async def show_main_menu(
    update: Update, context: ContextTypes.DEFAULT_TYPE, text: str = None
) -> None:
    """Show the main menu to the user with current Iran time."""
    iran_time = format_iran_time()
    header = f"🕐 <b>{iran_time}</b>\n\n"
    msg = text or (header + MESSAGES["main_menu"])
    markup = get_main_menu_keyboard()

    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(
            msg, reply_markup=markup, parse_mode="HTML"
        )
    else:
        await update.message.reply_text(
            msg, reply_markup=markup, parse_mode="HTML"
        )


async def handle_menu_callbacks(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """
    Handle menu callbacks that DO NOT start conversations.
    Conversation starters (addtask, deletetask, complete, edittask)
    are handled by the ConversationHandler entry_points instead.
    """
    query = update.callback_query
    await query.answer()
    data = query.data
    uid = update.effective_user.id

    if data == "main_tasks":
        tasks = db.get_all_user_tasks(uid)
        msg = format_task_list(tasks, "📋 <b>لیست تمام وظایف:</b>")
        await query.edit_message_text(msg, parse_mode="HTML", reply_markup=get_back_button(),
                                      disable_web_page_preview=True)

    elif data == "main_today":
        tasks = db.get_today_tasks(uid)
        msg = format_task_list(tasks, "📅 <b>وظایف امروز:</b>")
        await query.edit_message_text(msg, parse_mode="HTML", reply_markup=get_back_button(),
                                      disable_web_page_preview=True)

    elif data == "main_overdue":
        tasks = db.get_overdue_tasks(uid)
        if not tasks:
            msg = MESSAGES["no_overdue_tasks"]
        else:
            msg = format_task_list(tasks, "⚠️ <b>وظایف عقب‌افتاده:</b>")
        await query.edit_message_text(msg, parse_mode="HTML", reply_markup=get_back_button(),
                                      disable_web_page_preview=True)

    elif data == "main_summary":
        today = get_today_shamsi()
        al = db.get_all_user_tasks(uid)
        td = db.get_today_tasks(uid)
        ov = db.get_overdue_tasks(uid)
        pe = db.get_user_tasks(uid, status="pending")
        co = db.get_user_tasks(uid, status="completed")
        msg = (
            f"📊 <b>خلاصه وظایف</b>\n\n📅 {today}\n\n"
            f"📋 کل: {len(al)}\n⏳ در انتظار: {len(pe)}\n"
            f"✅ تکمیل: {len(co)}\n📅 امروز: {len(td)}\n⚠️ عقب‌افتاده: {len(ov)}"
        )
        await query.edit_message_text(msg, parse_mode="HTML", reply_markup=get_back_button())

    elif data == "main_help":
        await query.edit_message_text(
            MESSAGES["help"], parse_mode="HTML", reply_markup=get_back_button()
        )

    elif data == "main_export":
        tasks = db.get_all_user_tasks(uid)
        if not tasks:
            await query.edit_message_text(MESSAGES["no_tasks"], reply_markup=get_back_button())
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Tasks"
            headers = ["ID", "عنوان", "توضیحات", "تاریخ", "زمان", "یادآوری", "وضعیت"]
            ws.append(headers)
            hf = Font(bold=True, color="FFFFFF")
            hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            for c in ws[1]:
                c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center")
            for t in tasks:
                rem = f"{t['reminder_minutes']} دقیقه قبل" if t['reminder_minutes'] else "بدون"
                st = STATUS_EMOJIS.get(t['status'], '') + " " + t['status']
                ws.append([t['id'], t['title'], t.get('description', ''), t['due_date'], t['due_time'], rem, st])
            for col in ws.columns:
                ml = max((len(str(c.value or '')) for c in col), default=0)
                ws.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)
            await query.edit_message_text("📤 در حال ارسال...")
            await update.effective_user.send_document(
                document=buf,
                filename=f"tasks_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                caption=MESSAGES["export_success"],
            )
            await update.effective_user.send_message(
                MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML"
            )
        except Exception as e:
            logger.error(f"Export: {e}")
            await query.edit_message_text(f"❌ خطا: {str(e)}", reply_markup=get_back_button())

    elif data == "main_import":
        await query.edit_message_text(
            "📥 یک فایل اکسل (xlsx.) آپلود کنید.\n"
            "ستون‌ها: عنوان | توضیحات | تاریخ | زمان | یادآوری\n"
            "ردیف اول هدر است.",
            reply_markup=get_back_button(),
        )

    elif data == "main_menu":
        await show_main_menu(update, context)


# -----------------------------------------------------------------------------
# Validation Helpers
# -----------------------------------------------------------------------------
def validate_date(date_str: str) -> Optional[str]:
    try:
        datetime.strptime(date_str, "%Y-%m-%d")
        return date_str
    except ValueError:
        return None


def validate_time(time_str: str) -> Optional[str]:
    try:
        datetime.strptime(time_str, "%H:%M")
        return time_str
    except ValueError:
        return None


def parse_reminder_option(text: str) -> Optional[int]:
    mapping = {"5 دقیقه قبل": 5, "15 دقیقه قبل": 15, "1 ساعت قبل": 60, "1 روز قبل": 1440}
    return mapping.get(text.strip())


def parse_field_option(text: str) -> Optional[str]:
    mapping = {"عنوان": "title", "توضیحات": "description", "تاریخ": "due_date",
               "زمان": "due_time", "یادآوری": "reminder_minutes"}
    return mapping.get(text.strip())


def format_task_list(tasks: List[Dict[str, Any]], title: str) -> str:
    if not tasks:
        return MESSAGES["no_tasks"]
    msg = f"{title}\n\n"
    for i, t in enumerate(tasks, 1):
        em = STATUS_EMOJIS.get(t["status"], "⏳")
        msg += f"━━━━━━━━━━━━━━━━━━\n{i}. {em} <b>{t['title']}</b>\n"
        if t["description"]:
            msg += f"   📝 {t['description']}\n"
        msg += f"   📅 {t['due_date']} ⏰ {t['due_time']}\n"
        if t["reminder_minutes"]:
            msg += f"   🔔 {t['reminder_minutes']} دقیقه قبل\n"
    return msg + "━━━━━━━━━━━━━━━━━━\n"


def format_single_task(t: Dict[str, Any]) -> str:
    em = STATUS_EMOJIS.get(t["status"], "⏳")
    return (
        f"━━━━━━━━━━━━━━━━━━\n{em} <b>{t['title']}</b>\n━━━━━━━━━━━━━━━━━━\n"
        f"🆔 <code>{t['id']}</code>\n📝 {t.get('description', '—')}\n"
        f"📅 {t['due_date']}\n⏰ {t['due_time']}\n"
        f"🔔 {t['reminder_minutes']} دقیقه قبل\n📊 {em} {t['status']}\n━━━━━━━━━━━━━━━━━━\n"
    )


# -----------------------------------------------------------------------------
# Command Handlers
# -----------------------------------------------------------------------------
async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    u = update.effective_user
    db.add_or_update_user(u.id, u.username, u.first_name, u.last_name)
    await update.message.reply_text(
        f"👋 <b>خوش آمدید {u.first_name} عزیز!</b>\n\n"
        f"🤖 ربات مدیریت وظایف\n✅ از منو استفاده کنید.",
        reply_markup=get_main_menu_keyboard(), parse_mode="HTML",
    )


async def cmd_help(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await update.message.reply_text(MESSAGES["help"], parse_mode="HTML", reply_markup=get_back_button())


async def cmd_menu(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    await show_main_menu(update, context)


async def cmd_tasks(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tasks = db.get_all_user_tasks(update.effective_user.id)
    await update.message.reply_text(
        format_task_list(tasks, "📋 <b>لیست تمام وظایف:</b>"),
        parse_mode="HTML", reply_markup=get_back_button(), disable_web_page_preview=True,
    )


async def cmd_today(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    tasks = db.get_today_tasks(update.effective_user.id)
    await update.message.reply_text(
        format_task_list(tasks, "📅 <b>وظایف امروز:</b>"),
        parse_mode="HTML", reply_markup=get_back_button(), disable_web_page_preview=True,
    )


async def cmd_overdue(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    uid = update.effective_user.id
    tasks = db.get_overdue_tasks(uid)
    if not tasks:
        await update.message.reply_text(MESSAGES["no_overdue_tasks"], reply_markup=get_back_button())
        return
    await update.message.reply_text(
        format_task_list(tasks, "⚠️ <b>وظایف عقب‌افتاده:</b>"),
        parse_mode="HTML", reply_markup=get_back_button(), disable_web_page_preview=True,
    )


async def cmd_summary(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    uid = update.effective_user.id
    today = get_today_shamsi()
    al, td, ov = db.get_all_user_tasks(uid), db.get_today_tasks(uid), db.get_overdue_tasks(uid)
    pe, co = db.get_user_tasks(uid, status="pending"), db.get_user_tasks(uid, status="completed")
    msg = (f"📊 <b>خلاصه وظایف</b>\n\n📅 {today}\n\n━━━━━━━━━━━━━━━━━━\n"
           f"📋 کل: {len(al)}\n⏳ در انتظار: {len(pe)}\n✅ تکمیل: {len(co)}\n"
           f"📅 امروز: {len(td)}\n⚠️ عقب‌افتاده: {len(ov)}\n━━━━━━━━━━━━━━━━━━")
    await update.message.reply_text(msg, parse_mode="HTML", reply_markup=get_back_button())


# -----------------------------------------------------------------------------
# Add Task Conversation
# -----------------------------------------------------------------------------
ADD_TITLE, ADD_DESC, ADD_DATE, ADD_TIME, ADD_REM = range(5)


async def addtask_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    txt = MESSAGES["enter_title"]
    mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(txt, reply_markup=mk)
    else:
        await update.message.reply_text(txt, reply_markup=mk)
    return ADD_TITLE


async def addtask_title(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data["t"] = update.message.text.strip()
    mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
    await update.message.reply_text(MESSAGES["enter_description"], reply_markup=mk)
    return ADD_DESC


async def addtask_desc(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data["d"] = update.message.text.strip()
    mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
    await update.message.reply_text(MESSAGES["enter_due_date"], reply_markup=mk)
    return ADD_DATE


async def addtask_date(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    ds = update.message.text.strip()
    if not validate_date(ds):
        mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
        await update.message.reply_text(MESSAGES["invalid_date"], reply_markup=mk)
        return ADD_DATE
    context.user_data["dd"] = ds
    mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
    await update.message.reply_text(MESSAGES["enter_due_time"], reply_markup=mk)
    return ADD_TIME


async def addtask_time(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    ts = update.message.text.strip()
    if not validate_time(ts):
        mk = InlineKeyboardMarkup([[InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")]])
        await update.message.reply_text(MESSAGES["invalid_time"], reply_markup=mk)
        return ADD_TIME
    context.user_data["dt"] = ts
    mk = InlineKeyboardMarkup([
        [InlineKeyboardButton("🕐 5 دقیقه", callback_data="rem_5"),
         InlineKeyboardButton("🕐 15 دقیقه", callback_data="rem_15")],
        [InlineKeyboardButton("🕐 1 ساعت", callback_data="rem_60"),
         InlineKeyboardButton("🕐 1 روز", callback_data="rem_1440")],
        [InlineKeyboardButton("❌ بدون یادآوری", callback_data="rem_0")],
        [InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")],
    ])
    await update.message.reply_text(MESSAGES["select_reminder"], reply_markup=mk)
    return ADD_REM


async def addtask_rem(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    rm = int(q.data.replace("rem_", ""))
    uid = update.effective_user.id
    tid = db.add_task(uid, context.user_data.get("t", "?"), context.user_data.get("d", ""),
                      context.user_data.get("dd", "2000-01-01"), context.user_data.get("dt", "00:00"), rm)
    if rm > 0:
        scheduler.schedule_reminder(tid, uid, context.user_data["dd"], context.user_data["dt"], rm, send_reminder)
    context.user_data.clear()
    await q.edit_message_text(f"✅ <b>وظیفه با موفقیت اضافه شد!</b>", parse_mode="HTML")
    await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


# -----------------------------------------------------------------------------
# Delete Task Conversation
# -----------------------------------------------------------------------------
DEL_SELECT = range(1)


async def deletetask_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    uid = update.effective_user.id
    tasks = db.get_all_user_tasks(uid)
    if not tasks:
        txt = MESSAGES["no_tasks"]
        if update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.edit_message_text(txt)
        else:
            await update.message.reply_text(txt)
        await _back_to_menu(update)
        return ConversationHandler.END

    kb = []
    for t in tasks:
        em = STATUS_EMOJIS.get(t["status"], "⏳")
        kb.append([InlineKeyboardButton(f"{em} {t['title']} ({t['due_date']})", callback_data=f"del_{t['id']}")])
    kb.append([InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")])
    mk = InlineKeyboardMarkup(kb)
    txt = "🗑 <b>وظیفه را انتخاب کنید:</b>"

    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(txt, reply_markup=mk, parse_mode="HTML")
    else:
        await update.message.reply_text(txt, reply_markup=mk, parse_mode="HTML")
    return DEL_SELECT


async def deletetask_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    data = q.data

    if data in ("exit_conv",):
        await q.edit_message_text(MESSAGES["cancelled"])
        await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
        return ConversationHandler.END

    tid = int(data.replace("del_", ""))
    uid = update.effective_user.id
    scheduler.remove_task_reminder(tid, uid)
    db.delete_task(tid)

    await q.edit_message_text("🗑 <b>حذف شد.</b>", parse_mode="HTML")

    remaining = db.get_all_user_tasks(uid)
    if remaining:
        kb = []
        for t in remaining:
            em = STATUS_EMOJIS.get(t["status"], "⏳")
            kb.append([InlineKeyboardButton(f"{em} {t['title']}", callback_data=f"del_{t['id']}")])
        kb.append([InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")])
        await q.message.reply_text("🗑 <b>وظایف باقی‌مانده:</b>", reply_markup=InlineKeyboardMarkup(kb), parse_mode="HTML")
        return DEL_SELECT

    await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


# -----------------------------------------------------------------------------
# Complete Task Conversation
# -----------------------------------------------------------------------------
COMP_SELECT = range(1)


async def completetask_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    uid = update.effective_user.id
    tasks = db.get_user_tasks(uid, status="pending")
    if not tasks:
        txt = MESSAGES["no_tasks"]
        if update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.edit_message_text(txt)
        else:
            await update.message.reply_text(txt)
        await _back_to_menu(update)
        return ConversationHandler.END

    kb = [[InlineKeyboardButton(f"⏳ {t['title']} ({t['due_date']})", callback_data=f"comp_{t['id']}")] for t in tasks]
    kb.append([InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")])
    mk = InlineKeyboardMarkup(kb)
    txt = "✅ <b>وظیفه را انتخاب کنید:</b>"

    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(txt, reply_markup=mk, parse_mode="HTML")
    else:
        await update.message.reply_text(txt, reply_markup=mk, parse_mode="HTML")
    return COMP_SELECT


async def completetask_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    data = q.data

    if data == "exit_conv":
        await q.edit_message_text(MESSAGES["cancelled"])
        await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
        return ConversationHandler.END

    tid = int(data.replace("comp_", ""))
    uid = update.effective_user.id
    db.mark_task_completed(tid)
    scheduler.remove_task_reminder(tid, uid)

    await q.edit_message_text("✅ <b>تکمیل شد! 🎉</b>", parse_mode="HTML")

    remaining = db.get_user_tasks(uid, status="pending")
    if remaining:
        kb = [[InlineKeyboardButton(f"⏳ {t['title']}", callback_data=f"comp_{t['id']}")] for t in remaining]
        kb.append([InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")])
        await q.message.reply_text("📋 <b>باقی‌مانده:</b>", reply_markup=InlineKeyboardMarkup(kb), parse_mode="HTML")
        return COMP_SELECT

    await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


# -----------------------------------------------------------------------------
# Edit Task Conversation
# -----------------------------------------------------------------------------
ED_SELECT, ED_FIELD, ED_TITLE, ED_DESC, ED_DATE, ED_TIME, ED_REM = range(7)


async def edittask_start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    uid = update.effective_user.id
    tasks = db.get_all_user_tasks(uid)
    if not tasks:
        txt = MESSAGES["no_tasks"]
        if update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.edit_message_text(txt)
        else:
            await update.message.reply_text(txt)
        await _back_to_menu(update)
        return ConversationHandler.END

    kb = []
    for t in tasks:
        em = STATUS_EMOJIS.get(t["status"], "⏳")
        kb.append([InlineKeyboardButton(f"{em} {t['title']} ({t['due_date']})", callback_data=f"ed_{t['id']}")])
    kb.append([InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")])
    mk = InlineKeyboardMarkup(kb)
    txt = "✏️ <b>وظیفه را انتخاب کنید:</b>"

    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(txt, reply_markup=mk, parse_mode="HTML")
    else:
        await update.message.reply_text(txt, reply_markup=mk, parse_mode="HTML")
    return ED_SELECT


async def edittask_select(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    data = q.data
    if data == "exit_conv":
        await q.edit_message_text(MESSAGES["cancelled"])
        await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
        return ConversationHandler.END

    tid = int(data.replace("ed_", ""))
    context.user_data["eid"] = tid
    context.user_data["euid"] = update.effective_user.id
    task = db.get_task(tid)
    if not task:
        await q.edit_message_text("❌ یافت نشد.")
        return ConversationHandler.END

    info = format_single_task(task)
    kb = InlineKeyboardMarkup([
        [InlineKeyboardButton("📝 عنوان", callback_data="f_title"),
         InlineKeyboardButton("📄 توضیحات", callback_data="f_desc")],
        [InlineKeyboardButton("📅 تاریخ", callback_data="f_date"),
         InlineKeyboardButton("⏰ زمان", callback_data="f_time")],
        [InlineKeyboardButton("🔔 یادآوری", callback_data="f_rem")],
        [InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")],
    ])
    await q.edit_message_text(f"{info}\n\n📝 <b>کدام فیلد؟</b>", reply_markup=kb, parse_mode="HTML")
    return ED_FIELD


async def edittask_field(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    data = q.data
    if data == "exit_conv":
        await q.edit_message_text(MESSAGES["cancelled"])
        await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
        return ConversationHandler.END

    mapping = {
        "f_title": ("title", ED_TITLE), "f_desc": ("description", ED_DESC),
        "f_date": ("due_date", ED_DATE), "f_time": ("due_time", ED_TIME),
        "f_rem": ("reminder_minutes", ED_REM),
    }
    fld, nxt = mapping.get(data, (None, None))
    if fld is None:
        await q.edit_message_text("❌ نامعتبر.")
        return ConversationHandler.END

    context.user_data["ef"] = fld
    if fld == "reminder_minutes":
        mk = InlineKeyboardMarkup([
            [InlineKeyboardButton("🕐 5 دقیقه", callback_data="er_5"),
             InlineKeyboardButton("🕐 15 دقیقه", callback_data="er_15")],
            [InlineKeyboardButton("🕐 1 ساعت", callback_data="er_60"),
             InlineKeyboardButton("🕐 1 روز", callback_data="er_1440")],
            [InlineKeyboardButton("❌ بدون", callback_data="er_0")],
            [InlineKeyboardButton("🏠 منوی اصلی", callback_data="exit_conv")],
        ])
        await q.edit_message_text(MESSAGES["select_reminder"], reply_markup=mk)
        return ED_REM
    await q.edit_message_text("✏️ مقدار جدید را وارد کنید:")
    return nxt


async def edittask_title(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    v = update.message.text.strip()
    db.update_task(context.user_data["eid"], title=v)
    _resched(context.user_data["eid"], context)
    context.user_data.clear()
    await update.message.reply_text("✅ <b>به‌روزرسانی شد.</b>", parse_mode="HTML")
    await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


async def edittask_desc(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    v = update.message.text.strip()
    db.update_task(context.user_data["eid"], description=v)
    context.user_data.clear()
    await update.message.reply_text("✅ <b>به‌روزرسانی شد.</b>", parse_mode="HTML")
    await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


async def edittask_date(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    v = update.message.text.strip()
    if not validate_date(v):
        await update.message.reply_text(MESSAGES["invalid_date"])
        return ED_DATE
    db.update_task(context.user_data["eid"], due_date=v)
    _resched(context.user_data["eid"], context)
    context.user_data.clear()
    await update.message.reply_text("✅ <b>به‌روزرسانی شد.</b>", parse_mode="HTML")
    await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


async def edittask_time(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    v = update.message.text.strip()
    if not validate_time(v):
        await update.message.reply_text(MESSAGES["invalid_time"])
        return ED_TIME
    db.update_task(context.user_data["eid"], due_time=v)
    _resched(context.user_data["eid"], context)
    context.user_data.clear()
    await update.message.reply_text("✅ <b>به‌روزرسانی شد.</b>", parse_mode="HTML")
    await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


async def edittask_rem(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    data = q.data
    if data == "exit_conv":
        await q.edit_message_text(MESSAGES["cancelled"])
        await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
        return ConversationHandler.END

    rm = int(data.replace("er_", ""))
    tid = context.user_data["eid"]
    uid = update.effective_user.id
    db.update_task(tid, reminder_minutes=rm)
    scheduler.remove_task_reminder(tid, uid)
    if rm > 0:
        t = db.get_task(tid)
        if t:
            scheduler.schedule_reminder(tid, uid, t["due_date"], t["due_time"], rm, send_reminder)
    context.user_data.clear()
    await q.edit_message_text("✅ <b>به‌روزرسانی شد.</b>", parse_mode="HTML")
    await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


def _resched(tid: int, ctx: ContextTypes.DEFAULT_TYPE) -> None:
    u = ctx.user_data.get("euid")
    if not u:
        t = db.get_task(tid)
        if t:
            u = t["user_id"]
    if u:
        t = db.get_task(tid)
        if t and t["reminder_minutes"] > 0:
            scheduler.remove_task_reminder(tid, u)
            scheduler.schedule_reminder(tid, u, t["due_date"], t["due_time"], t["reminder_minutes"], send_reminder)


async def _back_to_menu(update: Update) -> None:
    if update.callback_query:
        await update.callback_query.message.reply_text(
            MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    else:
        await update.message.reply_text(
            MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")


# -----------------------------------------------------------------------------
# Export / Import
# -----------------------------------------------------------------------------
async def cmd_export(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    uid = update.effective_user.id
    tasks = db.get_all_user_tasks(uid)
    if not tasks:
        await update.message.reply_text(MESSAGES["no_tasks"], reply_markup=get_back_button())
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.append(["ID", "عنوان", "توضیحات", "تاریخ", "زمان", "یادآوری", "وضعیت"])
        hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        for c in ws[1]:
            c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center")
        for t in tasks:
            rem = f"{t['reminder_minutes']} دقیقه" if t['reminder_minutes'] else "بدون"
            st = STATUS_EMOJIS.get(t['status'], '') + " " + t['status']
            ws.append([t['id'], t['title'], t.get('description', ''), t['due_date'], t['due_time'], rem, st])
        for col in ws.columns:
            ml = max((len(str(c.value or '')) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        await update.message.reply_document(document=buf, filename=f"tasks_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                                            caption=MESSAGES["export_success"])
        await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    except Exception as e:
        logger.error(f"Export: {e}")
        await update.message.reply_text(f"❌ خطا: {str(e)}")


async def cmd_import(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    uid = update.effective_user.id
    if not update.message.document:
        await update.message.reply_text(MESSAGES["import_error"])
        return
    try:
        import openpyxl
        f = update.message.document
        if not f.file_name.endswith(".xlsx"):
            await update.message.reply_text(MESSAGES["import_error"])
            return
        fb = await f.get_file()
        fc = await fb.download_as_bytearray()
        wb = openpyxl.load_workbook(io.BytesIO(fc))
        ws = wb.active
        cnt = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            title, desc, dd, dt = str(row[1] or ""), str(row[2] or ""), str(row[3] or ""), str(row[4] or "")
            if not title or not dd or not dt:
                continue
            rm = 0
            rt = str(row[5] or "0")
            if "5" in rt: rm = 5
            elif "15" in rt: rm = 15
            elif "60" in rt or "ساعت" in rt: rm = 60
            elif "1440" in rt or "روز" in rt: rm = 1440
            tid = db.add_task(uid, title, desc, dd, dt, rm)
            if rm > 0:
                scheduler.schedule_reminder(tid, uid, dd, dt, rm, send_reminder)
            cnt += 1
        await update.message.reply_text(f"✅ <b>{MESSAGES['import_success']}</b>\n📊 {cnt} وظیفه.", parse_mode="HTML")
        await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    except ImportError:
        await update.message.reply_text("❌ openpyxl نصب نیست.\npip install openpyxl")
    except Exception as e:
        logger.error(f"Import: {e}")
        await update.message.reply_text(f"❌ خطا: {str(e)}")


# -----------------------------------------------------------------------------
# Daily Summary
# -----------------------------------------------------------------------------
async def send_daily_summary() -> None:
    try:
        from telegram import Bot
        bot = Bot(token=BOT_TOKEN)
        users = db.get_all_users()
        today = get_today_shamsi()
        for u in users:
            uid = u["user_id"]
            try:
                td = db.get_today_tasks(uid); ov = db.get_overdue_tasks(uid); pe = db.get_user_tasks(uid, status="pending")
                msg = (f"🌅 <b>صبح بخیر!</b>\n\n📊 <b>خلاصه روزانه</b>\n📅 {today}\n\n━━━━━━━━━━━━━━━━━━\n"
                       f"⏳ در انتظار: {len(pe)}\n📅 امروز: {len(td)}\n⚠️ عقب‌افتاده: {len(ov)}\n"
                       f"━━━━━━━━━━━━━━━━━━\n\n")
                if td:
                    msg += "📋 <b>امروز:</b>\n"
                    for t in td[:5]:
                        msg += f"• {STATUS_EMOJIS.get(t['status'],'⏳')} {t['title']} ⏰ {t['due_time']}\n"
                    if len(td) > 5:
                        msg += f"... و {len(td)-5} تا دیگر\n"
                    msg += "\n"
                msg += MESSAGES["main_menu"]
                await bot.send_message(chat_id=uid, text=msg, parse_mode="HTML", reply_markup=get_main_menu_keyboard())
            except Exception as e:
                logger.error(f"Summary user {uid}: {e}")
    except Exception as e:
        logger.error(f"Summary: {e}")


# -----------------------------------------------------------------------------
# Reminder
# -----------------------------------------------------------------------------
async def send_reminder(task_id: int, user_id: int) -> None:
    try:
        from telegram import Bot
        bot = Bot(token=BOT_TOKEN)
        t = db.get_task(task_id)
        if not t:
            return
        msg = (f"🔔 <b>یادآوری!</b>\n\n━━━━━━━━━━━━━━━━━━\n📌 <b>{t['title']}</b>\n"
               f"━━━━━━━━━━━━━━━━━━\n📝 {t.get('description','—')}\n📅 {t['due_date']}\n⏰ {t['due_time']}\n"
               f"━━━━━━━━━━━━━━━━━━")
        await bot.send_message(chat_id=user_id, text=msg, parse_mode="HTML", reply_markup=get_main_menu_keyboard())
    except Exception as e:
        logger.error(f"Reminder task {task_id}: {e}")


# -----------------------------------------------------------------------------
# Cancel
# -----------------------------------------------------------------------------
async def cmd_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    await update.message.reply_text(MESSAGES["cancelled"])
    await update.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


async def cancel_conv(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    q = update.callback_query
    await q.answer()
    context.user_data.clear()
    await q.edit_message_text(MESSAGES["cancelled"])
    await q.message.reply_text(MESSAGES["main_menu"], reply_markup=get_main_menu_keyboard(), parse_mode="HTML")
    return ConversationHandler.END


# -----------------------------------------------------------------------------
# Error Handler
# -----------------------------------------------------------------------------
async def error_handler(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    logger.error(f"Update {update} caused error {context.error}")
    try:
        if update and update.effective_message:
            await update.effective_message.reply_text("❌ خطایی رخ داد. دوباره تلاش کنید.")
    except Exception:
        pass


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def main() -> None:
    app = Application.builder().token(BOT_TOKEN).build()

    # Simple commands
    for cmd, fn in [("start", cmd_start), ("menu", cmd_menu), ("help", cmd_help),
                    ("tasks", cmd_tasks), ("today", cmd_today), ("overdue", cmd_overdue),
                    ("summary", cmd_summary), ("export", cmd_export), ("import", cmd_import)]:
        app.add_handler(CommandHandler(cmd, fn))

    # --- Conversation: Add Task ---
    app.add_handler(ConversationHandler(
        entry_points=[CommandHandler("addtask", addtask_start),
                      CallbackQueryHandler(addtask_start, pattern="^main_addtask$")],
        states={
            ADD_TITLE:  [MessageHandler(filters.TEXT & ~filters.COMMAND, addtask_title)],
            ADD_DESC:   [MessageHandler(filters.TEXT & ~filters.COMMAND, addtask_desc)],
            ADD_DATE:   [MessageHandler(filters.TEXT & ~filters.COMMAND, addtask_date)],
            ADD_TIME:   [MessageHandler(filters.TEXT & ~filters.COMMAND, addtask_time)],
            ADD_REM:    [CallbackQueryHandler(addtask_rem, pattern="^rem_"),
                         CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
        },
        fallbacks=[CommandHandler("cancel", cmd_cancel),
                   CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
    ))

    # --- Conversation: Delete Task ---
    app.add_handler(ConversationHandler(
        entry_points=[CommandHandler("deletetask", deletetask_start),
                      CallbackQueryHandler(deletetask_start, pattern="^main_deletetask$")],
        states={
            DEL_SELECT: [CallbackQueryHandler(deletetask_confirm, pattern="^del_"),
                         CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
        },
        fallbacks=[CommandHandler("cancel", cmd_cancel),
                   CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
    ))

    # --- Conversation: Complete Task ---
    app.add_handler(ConversationHandler(
        entry_points=[CommandHandler("complete", completetask_start),
                      CallbackQueryHandler(completetask_start, pattern="^main_complete$")],
        states={
            COMP_SELECT: [CallbackQueryHandler(completetask_confirm, pattern="^comp_"),
                          CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
        },
        fallbacks=[CommandHandler("cancel", cmd_cancel),
                   CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
    ))

    # --- Conversation: Edit Task ---
    app.add_handler(ConversationHandler(
        entry_points=[CommandHandler("edittask", edittask_start),
                      CallbackQueryHandler(edittask_start, pattern="^main_edittask$")],
        states={
            ED_SELECT: [CallbackQueryHandler(edittask_select, pattern="^ed_"),
                        CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
            ED_FIELD:  [CallbackQueryHandler(edittask_field, pattern="^f_"),
                        CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
            ED_TITLE:  [MessageHandler(filters.TEXT & ~filters.COMMAND, edittask_title)],
            ED_DESC:   [MessageHandler(filters.TEXT & ~filters.COMMAND, edittask_desc)],
            ED_DATE:   [MessageHandler(filters.TEXT & ~filters.COMMAND, edittask_date)],
            ED_TIME:   [MessageHandler(filters.TEXT & ~filters.COMMAND, edittask_time)],
            ED_REM:    [CallbackQueryHandler(edittask_rem, pattern="^er_"),
                        CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
        },
        fallbacks=[CommandHandler("cancel", cmd_cancel),
                   CallbackQueryHandler(cancel_conv, pattern="^exit_conv$")],
    ))

    # Menu callback handler (for non-conversation actions) - registered AFTER conversations
    app.add_handler(CallbackQueryHandler(handle_menu_callbacks, pattern="^main_"))
    app.add_handler(CallbackQueryHandler(show_main_menu, pattern="^main_menu$"))

    app.add_error_handler(error_handler)

    scheduler.schedule_daily_summary(send_daily_summary, DAILY_SUMMARY_HOUR, DAILY_SUMMARY_MINUTE)
    logger.info("Starting bot...")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()